import io
import unittest
from datetime import date
from unittest.mock import Mock, patch
from openpyxl import load_workbook
from app import app, fetch_reports, normalize_report, photo_urls


def record(**changes):
    row = dict(tanggal='2026-09-05', area_kerja='Utility', unit_area='Chiller',
               pekerjaan='Pemeriksaan pompa', status='Selesai', pic='Budi',
               foto=['https://example.com/photo.jpg'])
    row.update(changes)
    return row


def envelope():
    return dict(id='r1', created_date='2026-09-04T18:30:00+00:00', area_kerja='Utility', pic='Budi',
                report_items=[{'Unit/Area': 'Chiller', 'Pekerjaan': 'Pemeriksaan pompa',
                               'Status': 'Selesai', 'Foto': ['https://example.com/photo.jpg']},
                              {'unit': 'AHU', 'description': 'Cek filter', 'status': 'Proses', 'photos': []}])


def response(rows):
    return Mock(ok=True, status_code=200, json=lambda: rows)


class ReportTests(unittest.TestCase):
    def setUp(self):
        self.identity = patch('auth.get_identity', return_value={'token': 'test-token', 'email': 'test@example.com'})
        self.identity.start()
        self.addCleanup(self.identity.stop)
        self.client = app.test_client()

    def test_page_and_assets(self):
        for path in ('/', '/static/style.css', '/static/app.js'):
            result = self.client.get(path)
            self.assertEqual(result.status_code, 200)
            result.close()

    @patch('app.fetch_reports')
    def test_bad_period_does_not_fetch(self, fetch):
        for query in ('start=invalid', 'start=2026-09-05&end=2026-09-01'):
            self.assertEqual(self.client.get('/api/reports?' + query).status_code, 400)
            self.assertEqual(self.client.get('/export?' + query).status_code, 400)
        fetch.assert_not_called()

    @patch('app.fetch_reports')
    def test_combined_filters(self, fetch):
        fetch.return_value = [record(), record(pic='Ani', status='Proses'), record(area_kerja='Produksi')]
        data = self.client.get('/api/reports?area_kerja=Utility&unit_area=Chiller&status=Selesai&pic=Budi&q=POMPA').get_json()
        self.assertEqual(data['total'], 1)
        self.assertEqual(data['rows'][0]['pic'], 'Budi')
        self.assertEqual(data['options']['status'], ['Proses', 'Selesai'])

    @patch('app.requests.get')
    def test_pagination_and_inclusive_utc_end(self, get):
        get.side_effect = [response([envelope()]), response([envelope()]), response([])]
        result = fetch_reports(date(2026, 9, 1), date(2026, 9, 5), 'test-token')
        self.assertEqual(len(result), 4)
        params = get.call_args_list[0].kwargs['params']
        self.assertIn(('created_date', 'gte.2026-09-01T00:00:00+00:00'), params)
        self.assertIn(('created_date', 'lt.2026-09-06T00:00:00+00:00'), params)
        self.assertIn(('order', 'created_date.desc,id.desc'), params)
        self.assertIn(('offset', '2'), get.call_args_list[2].kwargs['params'])

    def test_flatten_items_and_inherit_parent_fields(self):
        rows = normalize_report(envelope())
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0], record(id='r1', tanggal='2026-09-04T18:30:00+00:00'))
        self.assertEqual(rows[1]['unit_area'], 'AHU')
        self.assertEqual(rows[1]['pekerjaan'], 'Cek filter')
        self.assertEqual(rows[1]['pic'], 'Budi')
        self.assertEqual(rows[1]['tanggal'], '2026-09-04T18:30:00+00:00')

    @patch('app.requests.get')
    def test_actual_schema_to_filtered_excel(self, get):
        get.side_effect = [response([envelope()]), response([])]
        result = self.client.get('/export?start=2026-09-01&end=2026-09-05&status=Proses')
        self.assertEqual(result.status_code, 200)
        book = load_workbook(io.BytesIO(result.data))
        self.assertEqual(book.active.max_row, 2)
        self.assertEqual(book.active['C2'].value, 'AHU')
        self.assertEqual(book.active['D2'].value, 'Cek filter')
        self.assertEqual(book.active['A2'].value.isoformat(), '2026-09-04T18:30:00')
        self.assertEqual(book.active['A2'].number_format, 'yyyy-mm-dd hh:mm:ss')
        book.close()

    @patch('app.fetch_reports')
    def test_excel_formula_safety(self, fetch):
        fetch.return_value = [record(pekerjaan='=HYPERLINK("evil")'), record(pic='Ani')]
        result = self.client.get('/export?pic=Budi')
        book = load_workbook(io.BytesIO(result.data))
        sheet = book.active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet['D2'].value, '=HYPERLINK("evil")')
        self.assertEqual(sheet['D2'].data_type, 's')
        self.assertEqual(sheet['A2'].value.date(), date(2026, 9, 5))
        self.assertEqual(sheet['G2'].hyperlink.target, 'https://example.com/photo.jpg')
        self.assertEqual(sheet.freeze_panes, 'A2')
        book.close()

    @patch('app.requests.get')
    def test_connection_error(self, get):
        get.return_value = Mock(status_code=401)
        result = self.client.get('/api/reports')
        self.assertEqual(result.status_code, 502)
        self.assertIn('menolak akses', result.get_json()['error'])

    @patch('app.requests.get')
    def test_empty_table_and_empty_items(self, get):
        get.return_value = response([])
        self.assertEqual(self.client.get('/api/reports').get_json()['total'], 0)
        row = envelope()
        row['report_items'] = []
        self.assertEqual(normalize_report(row), [])

    @patch('app.requests.get')
    def test_invalid_items(self, get):
        for items in (None, {}, [None], [{'unknown': 'value'}]):
            row = envelope()
            row['report_items'] = items
            get.side_effect = [response([row]), response([])]
            self.assertEqual(self.client.get('/api/reports').status_code, 503)

    def test_photo_urls(self):
        self.assertEqual(photo_urls(['javascript:alert(1)', 'https://example.com/a.png']), ['https://example.com/a.png'])


if __name__ == '__main__':
    unittest.main()
