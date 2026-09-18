import io
import json
import os
import re
import secrets
from pathlib import Path
from datetime import date, datetime, time, timedelta
from zoneinfo import ZoneInfo
from urllib.parse import urlparse

import requests
from dotenv import load_dotenv
from flask import Flask, jsonify, render_template, request, send_file, session, redirect, url_for, g
from flask_sqlalchemy import SQLAlchemy
import auth
from item_catalog import search_items
from pricelist_catalog import search_pricelist, format_price
from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PyPDF2 import PdfReader, PdfWriter
from sqlalchemy import text as sql_text
from sqlalchemy.orm import deferred
from werkzeug.utils import secure_filename

load_dotenv()
app = Flask(__name__)
app.config.update(SECRET_KEY=os.getenv('FLASK_SECRET_KEY') or secrets.token_hex(32),
                  AUTH_DB=str(Path(__file__).parent / 'instance' / 'sessions.sqlite3'),
                  SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Lax',
                  SESSION_COOKIE_SECURE=os.getenv('COOKIE_SECURE', 'false').lower() == 'true',
                  SESSION_PERMANENT=True,
                  PERMANENT_SESSION_LIFETIME=timedelta(days=3650),
                  SESSION_REFRESH_EACH_REQUEST=True,
                  MAX_CONTENT_LENGTH=50 * 1024 * 1024,
                  SQLALCHEMY_ENGINE_OPTIONS={'connect_args': {'timeout': 30}},
                  SQLALCHEMY_DATABASE_URI='sqlite:///berita_acara.db',
                  SQLALCHEMY_TRACK_MODIFICATIONS=False)

# Initialize SQLAlchemy
db = SQLAlchemy(app)


class PdfUpload(db.Model):
    __tablename__ = 'pdf_uploads'
    id = db.Column(db.Integer, primary_key=True)
    tanggal = db.Column(db.String(50), nullable=False)
    no_wpr = db.Column(db.String(50), nullable=False)
    deskripsi = db.Column(db.Text, nullable=False)
    type_budget = db.Column(db.String(50), nullable=False)
    keterangan = db.Column(db.Text, nullable=True)
    status = db.Column(db.String(30), nullable=True)
    pdf_filename = db.Column(db.String(255), nullable=True)
    pdf_content = deferred(db.Column(db.LargeBinary, nullable=True))
    created_at = db.Column(db.DateTime, default=datetime.utcnow)


class TenantUtility(db.Model):
    __tablename__ = 'tenant_utilities'
    id = db.Column(db.Integer, primary_key=True)
    unit = db.Column(db.String(100), nullable=False, index=True)
    month = db.Column(db.String(30), nullable=False, index=True)
    source_filename = db.Column(db.String(255), nullable=False)
    values_json = db.Column(db.Text, nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)


def key_name(value):
    return re.sub(r'[^a-z0-9]', '', value.casefold())


ITEM_ALIASES = {
    'unit_area': ('unit_area', 'unit', 'area', 'unit_name', 'location'),
    'pekerjaan': ('pekerjaan', 'description', 'job', 'work', 'deskripsi', 'activity'),
    'status': ('status',),
    'foto': ('foto', 'photos', 'photo', 'photo_url', 'photo_urls', 'images', 'image_url', 'imagepaths'),
}


FIELDS = ('tanggal', 'area_kerja', 'unit_area', 'pekerjaan', 'status', 'pic', 'foto')
HEADERS = ('Tanggal & Waktu (UTC+7)', 'Area Kerja', 'Unit/Area', 'Pekerjaan', 'Status', 'PIC', 'Foto')


def report_timezone():
    return ZoneInfo(os.getenv('REPORT_TIMEZONE', 'Asia/Jakarta'))


def report_timezone_label():
    tz = os.getenv('REPORT_TIMEZONE', 'Asia/Jakarta')
    if tz == 'Asia/Jakarta':
        return 'WIB'
    if tz == 'Asia/Makassar':
        return 'WITA'
    if tz == 'Asia/Jayapura':
        return 'WIT'
    return tz


def normalize_report(row):
    items = row.get('report_items')
    if not isinstance(items, list):
        raise ReportError('Kolom report_items harus berupa array JSON.', 503)
    try:
        # Prefer created_at from Supabase (system-generated UTC timestamp, immune to UPSERT overwrites).
        # Fall back to created_date if created_at is not present.
        raw_ts = row.get('created_at') or row.get('created_date')
        if not raw_ts:
            raise ValueError('Kolom tanggal kosong')
        timestamp_str = str(raw_ts).strip()
        if timestamp_str.endswith('Z'):
            timestamp_str = timestamp_str[:-1] + '+00:00'
        # Python 3.9 accepts only 3 or 6 fractional digits; Postgres can emit 1–6.
        timestamp_str = re.sub(r'(T\d{2}:\d{2}:\d{2})\.(\d+)',
                               lambda match: match[1] + '.' + match[2][:6].ljust(6, '0'), timestamp_str)
        
        timestamp = datetime.fromisoformat(timestamp_str)
        
        # If timestamp has no timezone, assume UTC (Supabase standard)
        if timestamp.tzinfo is None:
            timestamp = timestamp.replace(tzinfo=ZoneInfo('UTC'))
        
        # Convert UTC timestamp to target reporting timezone (WIB / Asia/Jakarta / UTC+7)
        report_datetime = timestamp.astimezone(report_timezone()).isoformat()
        
    except (ValueError, KeyError) as e:
        raise ReportError(f'created_at/created_date tidak valid pada laporan: {e}', 503)
    records = []
    for item in items:
        if not isinstance(item, dict):
            raise ReportError('Setiap report_items harus berupa objek JSON pekerjaan.', 503)
        keys = {key_name(k): v for k, v in item.items()}
        record = {'id': row.get('id'), 'tanggal': report_datetime, 'area_kerja': str(row.get('area_kerja') or ''),
                  'pic': pic_value(row.get('pic'))}
        recognized = False
        for field, aliases in ITEM_ALIASES.items():
            candidates = (os.getenv('COL_' + field.upper(), field),) + aliases
            value = None
            for candidate in candidates:
                if key_name(candidate) in keys:
                    value = keys[key_name(candidate)]
                    recognized = True
                    break
            if field == 'foto':
                record[field] = photo_urls(value)
            elif field == 'status':
                status_value = str(value) if value is not None else ''
                if status_value.startswith('ProgressStatus.'):
                    record[field] = status_value.replace('ProgressStatus.', '')
                elif status_value == 'ProgressStatus':
                    record[field] = ''
                else:
                    record[field] = status_value
            else:
                record[field] = str(value) if value is not None else ''
        if not recognized:
            raise ReportError('Key report_items belum cocok. Sesuaikan COL_UNIT_AREA, COL_PEKERJAAN, COL_STATUS, dan COL_FOTO di .env.', 503)
        records.append(record)
    return records


class ReportError(Exception):
    def __init__(self, message, status=502):
        super().__init__(message)
        self.status = status


def period(args):
    values = []
    for key in ('start', 'end'):
        value = args.get(key, '').strip()
        try:
            values.append(date.fromisoformat(value) if value else None)
        except ValueError:
            raise ReportError('Format tanggal harus YYYY-MM-DD.', 400)
    start, end = values
    if start and end and start > end:
        raise ReportError('Tanggal mulai tidak boleh melewati tanggal akhir.', 400)
    if end == date.max:
        raise ReportError('Tanggal akhir terlalu besar.', 400)
    return start, end


def photo_urls(value):
    if not value:
        return []
    if isinstance(value, dict):
        value = value.get('url') or value.get('publicUrl') or ''
    if isinstance(value, str):
        value = [value]
    if not isinstance(value, list):
        return []
    urls = []
    for item in value:
        if isinstance(item, dict):
            item = item.get('url', '')
        if isinstance(item, str) and urlparse(item).scheme in ('http', 'https') and urlparse(item).netloc:
            urls.append(item)
    return urls


def pic_value(value):
    if not value:
        return ''
    if isinstance(value, str):
        # If string contains spaces, assume it's a single full name
        # If it already has | separator, keep it
        if '|' in value:
            return value
        return value
    if isinstance(value, dict):
        return str(value.get('pic') or value.get('name') or value.get('email') or '')
    if isinstance(value, list):
        parts = []
        for item in value:
            if isinstance(item, dict):
                item = item.get('pic') or item.get('name') or item.get('email') or ''
            if item is not None and str(item).strip():
                parts.append(str(item))
        return '|'.join(parts)
    return str(value)


def fetch_reports(start, end, access_token, source='engineering'):
    prefix = 'DW_' if source == 'dw' else ''
    url = os.getenv(prefix + 'SUPABASE_URL', '').rstrip('/')
    key = os.getenv(prefix + 'SUPABASE_KEY', '')
    table = os.getenv(prefix + 'SUPABASE_TABLE', '').strip()
    if not table:
        raise ReportError(f'Nama tabel belum diatur. Isi {prefix}SUPABASE_TABLE pada file .env, lalu mulai ulang aplikasi.', 503)
    if not url or not key:
        raise ReportError(f'Isi {prefix}SUPABASE_URL dan {prefix}SUPABASE_KEY pada file .env.', 503)
    if not all(c.isalnum() or c == '_' for c in table):
        raise ReportError(f'{prefix}SUPABASE_TABLE harus berupa nama tabel, tanpa URL atau nama schema.', 503)
    # A login token from the Engineering project is not valid in the DW project.
    headers = ({'apikey': key, 'Authorization': 'Bearer ' + key, 'Accept': 'application/json'}
               if source == 'dw' else auth.api_headers(access_token))
    area_column = 'title' if source == 'dw' else 'area_kerja'
    params = [('select', f'id,created_at,created_date,pic,{area_column},report_items'),
              ('order', 'created_at.desc,id.desc')]
    if start:
        start_utc = datetime.combine(start, time.min, report_timezone()).astimezone(ZoneInfo('UTC')).isoformat()
        params.append(('created_at', 'gte.' + start_utc))
    if end:
        end_utc = datetime.combine(end + timedelta(days=1), time.min, report_timezone()).astimezone(ZoneInfo('UTC')).isoformat()
        params.append(('created_at', 'lt.' + end_utc))
    rows = []
    # Keep fetching even if the project has a server-side limit smaller than 1000.
    while True:
        try:
            response = requests.get(
                f'{url}/rest/v1/{table}',
                headers=headers,
                params=params + [('offset', str(len(rows))), ('limit', '1000')],
                timeout=(10, 30),
            )
        except requests.RequestException:
            raise ReportError('Tidak dapat terhubung ke Supabase. Periksa jaringan dan URL proyek.')
        if response.status_code in (401, 403):
            if source == 'dw':
                raise ReportError('Supabase DW menolak akses. Periksa DW_SUPABASE_KEY dan izin SELECT untuk role anon.', 403)
            raise ReportError('Supabase menolak akses. Periksa sesi login dan policy SELECT anggota tim untuk role authenticated.')
        if response.status_code in (400, 404):
            raise ReportError(f'Struktur tabel tidak cocok. Diperlukan kolom id, created_at, created_date, pic, {area_column}, dan report_items. Periksa {prefix}SUPABASE_TABLE pada .env.')
        if not response.ok:
            raise ReportError('Supabase sedang tidak tersedia. Silakan coba kembali.')
        try:
            batch = response.json()
        except ValueError:
            raise ReportError('Respons Supabase tidak valid.')
        if not isinstance(batch, list):
            raise ReportError('Respons Supabase bukan daftar laporan.')
        if not batch:
            break
        rows.extend(batch)
        if len(rows) > 100000:
            raise ReportError('Periode melebihi 100.000 laporan. Persempit rentang tanggal.', 400)
    normalized = []
    for row in rows:
        if source == 'dw':
            row = dict(row, area_kerja=row.get('title') or '')
        normalized.extend(normalize_report(row))
        if len(normalized) > 100000:
            raise ReportError('Periode melebihi 100.000 item pekerjaan. Persempit rentang tanggal.', 400)
    return normalized


def filtered_reports(args, source='engineering'):
    start, end = period(args)
    rows = (fetch_reports(start, end, g.identity['token'], source='dw') if source == 'dw'
            else fetch_reports(start, end, g.identity['token']))
    options = {field: sorted({r[field] for r in rows if r[field]}) for field in ('area_kerja', 'unit_area', 'status', 'pic')}
    for field in options:
        value = args.get(field, '').strip()
        if value:
            rows = [r for r in rows if r[field] == value]
    query = args.get('q', '').strip().casefold()
    if query:
        rows = [r for r in rows if query in ' '.join(r[f] for f in FIELDS if f != 'foto').casefold()]
    return rows, options, start, end


@app.before_request
def require_login():
    if request.endpoint in ('static', 'login') or request.endpoint is None:
        return
    if request.endpoint == 'logout':
        return
    g.identity = auth.get_identity()
    if not g.identity:
        if request.endpoint in ('reports', 'export', 'dw_reports', 'dw_export'):
            return jsonify(error='Sesi berakhir. Silakan login kembali.'), 401
        # For API endpoints, return JSON error instead of redirect
        if request.path.startswith('/api/'):
            return jsonify(error='Sesi berakhir. Silakan login kembali.'), 401
        return redirect(url_for('login'))


def csrf_valid():
    return bool(session.get('csrf')) and secrets.compare_digest(session['csrf'], request.form.get('csrf', ''))


@app.route('/login', methods=['GET', 'POST'])
def login():
    if 'csrf' not in session:
        session['csrf'] = secrets.token_urlsafe(32)
    error, status = None, 200
    if request.method == 'POST':
        if not csrf_valid():
            error, status = 'Form login kedaluwarsa. Muat ulang halaman lalu coba lagi.', 400
        else:
            email, password = request.form.get('email', '').strip(), request.form.get('password', '')
            if not email or not password:
                error, status = 'Isi email dan password.', 400
            else:
                try:
                    auth.login_user(email, password)
                    return redirect(url_for('index'))
                except auth.AuthError as exc:
                    error, status = str(exc), exc.status
    return render_template('login.html', error=error), status


@app.post('/logout')
def logout():
    if not csrf_valid():
        return jsonify(error='Form kedaluwarsa. Muat ulang halaman.'), 400
    auth.logout_user()
    return redirect(url_for('login'))


@app.errorhandler(auth.AuthError)
def auth_error(error):
    if request.endpoint in ('reports', 'export', 'dw_reports', 'dw_export') or request.path.startswith('/api/'):
        return jsonify(error=str(error)), error.status
    session.setdefault('csrf', secrets.token_urlsafe(32))
    return render_template('login.html', error=str(error)), error.status


@app.get('/')
def index():
    return render_template('index.html',
                           report_timezone=str(report_timezone()),
                           report_timezone_label=report_timezone_label())


@app.get('/laporan-dw')
def dw_index():
    return render_template('index.html', report_title='Laporan DW',
                           report_heading='Laporan DW',
                           report_description='Aktivitas tim, progres pekerjaan, dan dokumentasi DW',
                           reports_url=url_for('dw_reports'), export_url=url_for('dw_export'),
                           export_prefix='laporan_dw',
                           report_timezone=str(report_timezone()),
                           report_timezone_label=report_timezone_label())


@app.get('/api/dw/reports')
def dw_reports():
    rows, options, _, _ = filtered_reports(request.args, source='dw')
    return jsonify(rows=rows, options=options, total=len(rows))


@app.get('/kode-item')
def item_codes():
    query = request.args.get('q', '').strip()[:200]
    try:
        page = int(request.args.get('page', '1'))
    except ValueError:
        page = 1
    return render_template('items.html', query=query, **search_items(query, page))


@app.get('/pricelist')
def pricelist():
    query = request.args.get('q', '').strip()[:200]
    view = 'jasa' if request.args.get('view') == 'jasa' else 'detail'
    try:
        page = int(request.args.get('page', '1'))
    except ValueError:
        page = 1
    return render_template('pricelist.html', query=query, view=view, money=format_price,
                           **search_pricelist(query, request.args.get('category', ''), page, view))


@app.get('/api/reports')
def reports():
    rows, options, _, _ = filtered_reports(request.args)
    return jsonify(rows=rows, options=options, total=len(rows))


@app.get('/api/reports/<report_id>')
def report_detail(report_id):
    try:
        url = os.getenv('SUPABASE_URL', '').rstrip('/')
        key = os.getenv('SUPABASE_KEY', '')
        table = os.getenv('SUPABASE_TABLE', '').strip()
        
        if not table:
            raise ReportError('Nama tabel belum diatur. Isi SUPABASE_TABLE pada file .env.', 503)
        if not url or not key:
            raise ReportError('Isi SUPABASE_URL dan SUPABASE_KEY pada file .env.', 503)
            
        response = requests.get(
            f'{url}/rest/v1/{table}',
            headers=auth.api_headers(g.identity['token']),
            params={'id': f'eq.{report_id}', 'limit': '1'},
            timeout=(10, 30),
        )
        
        if response.status_code in (401, 403):
            raise ReportError('Akses ditolak. Periksa sesi login dan policy.', 403)
        if not response.ok:
            raise ReportError('Gagal mengambil detail laporan.')
            
        data = response.json()
        if not data or not isinstance(data, list) or len(data) == 0:
            raise ReportError('Laporan tidak ditemukan.', 404)
            
        return jsonify(data[0])
        
    except requests.RequestException:
        raise ReportError('Tidak dapat terhubung ke Supabase.')


@app.put('/api/reports/<report_id>')
def update_report(report_id):
    try:
        url = os.getenv('SUPABASE_URL', '').rstrip('/')
        key = os.getenv('SUPABASE_KEY', '')
        table = os.getenv('SUPABASE_TABLE', '').strip()
        
        if not table:
            raise ReportError('Nama tabel belum diatur. Isi SUPABASE_TABLE pada file .env.', 503)
        if not url or not key:
            raise ReportError('Isi SUPABASE_URL dan SUPABASE_KEY pada file .env.', 503)
            
        data = request.get_json()
        if not data:
            raise ReportError('Data tidak valid.', 400)
            
        response = requests.patch(
            f'{url}/rest/v1/{table}',
            headers=auth.api_headers(g.identity['token']),
            params={'id': f'eq.{report_id}'},
            json=data,
            timeout=(10, 30),
        )
        
        if response.status_code in (401, 403):
            raise ReportError('Akses ditolak. Periksa sesi login dan policy.', 403)
        if not response.ok:
            raise ReportError('Gagal mengupdate laporan.')
            
        return jsonify(success=True, message='Laporan berhasil diupdate.')
        
    except requests.RequestException:
        raise ReportError('Tidak dapat terhubung ke Supabase.')


@app.delete('/api/reports/<report_id>')
def delete_report(report_id):
    try:
        url = os.getenv('SUPABASE_URL', '').rstrip('/')
        key = os.getenv('SUPABASE_KEY', '')
        table = os.getenv('SUPABASE_TABLE', '').strip()
        
        if not table:
            raise ReportError('Nama tabel belum diatur. Isi SUPABASE_TABLE pada file .env.', 503)
        if not url or not key:
            raise ReportError('Isi SUPABASE_URL dan SUPABASE_KEY pada file .env.', 503)
            
        response = requests.delete(
            f'{url}/rest/v1/{table}',
            headers=auth.api_headers(g.identity['token']),
            params={'id': f'eq.{report_id}'},
            timeout=(10, 30),
        )
        
        if response.status_code in (401, 403):
            raise ReportError('Akses ditolak. Periksa sesi login dan policy.', 403)
        if not response.ok:
            raise ReportError('Gagal menghapus laporan.')
            
        return jsonify(success=True, message='Laporan berhasil dihapus.')
        
    except requests.RequestException:
        raise ReportError('Tidak dapat terhubung ke Supabase.')


@app.get('/laporan-dw/export', endpoint='dw_export')
@app.get('/export')
def export():
    is_dw = request.endpoint == 'dw_export'
    rows, _, start, end = (filtered_reports(request.args, source='dw') if is_dw
                            else filtered_reports(request.args))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Laporan DW' if is_dw else 'Laporan Pekerjaan'
    sheet.append(HEADERS)
    for record in rows:
        sheet.append([('\n'.join(record[f]) if f == 'foto' else record[f]) for f in FIELDS])
        # Explicit strings prevent untrusted database values becoming Excel formulas.
        for cell in sheet[sheet.max_row]:
            cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        try:
            cell = sheet.cell(sheet.max_row, 1)
            # Parse datetime with timezone support
            dt = datetime.fromisoformat(record['tanggal'])
            if dt.tzinfo is not None:
                dt = dt.astimezone(report_timezone()).replace(tzinfo=None)
            cell.value = dt
            cell.number_format = 'yyyy-mm-dd hh:mm:ss'
        except (ValueError, TypeError):
            pass
        if record['foto']:
            sheet.cell(sheet.max_row, 7).hyperlink = record['foto'][0]
            sheet.cell(sheet.max_row, 7).style = 'Hyperlink'
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='163F36')
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(vertical='center')
    sheet.row_dimensions[1].height = 28
    for col, width in zip('ABCDEFG', (25, 25, 25, 85, 22, 20, 55)):
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f'{"laporan_dw" if is_dw else "laporan"}_{start or "semua"}_{end or "terbaru"}.xlsx'
    return send_file(output, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.get('/pdf-upload')
def pdf_upload():
    return render_template('pdf_upload.html')


@app.get('/pemakaian-listrik-air')
def tenant_utility():
    return render_template('tenant_utility.html')


@app.route('/api/tenant-utilities', methods=['GET', 'POST'])
def tenant_utility_api():
    if request.method == 'POST':
        if not csrf_valid():
            return jsonify(error='Form kedaluwarsa. Muat ulang halaman.'), 400
        upload = request.files.get('utility_file')
        if not upload or not upload.filename:
            return jsonify(error='Pilih file Excel terlebih dahulu.'), 400
        if not upload.filename.lower().endswith(('.xlsx', '.xlsm')):
            return jsonify(error='Hanya file Excel (.xlsx/.xlsm) yang diperbolehkan.'), 400
        try:
            workbook = load_workbook(upload, data_only=True, read_only=True)
            imported = 0
            month_names = {'januari','februari','maret','april','mei','juni','juli','agustus','september','oktober','november','desember',
                           'january','february','march','april','may','june','july','august','september','october','november','december'}
            month_map = {'januari':'Januari','februari':'Februari','maret':'Maret','april':'April','mei':'Mei','juni':'Juni','juli':'Juli','agustus':'Agustus','september':'September','oktober':'Oktober','november':'November','desember':'Desember','january':'Januari','february':'Februari','march':'Maret','may':'Mei','june':'Juni','july':'Juli','august':'Agustus','october':'Oktober','december':'Desember'}
            file_month = next((month_map[name] for name in month_map if name in upload.filename.casefold()), '')
            for sheet in workbook.worksheets:
                rows = [list(row) for row in sheet.iter_rows(values_only=True)]
                if not rows:
                    continue
                unit = ''
                for row in rows[:8]:
                    text = ' '.join(str(value or '') for value in row)
                    match = re.search(r'\b(?:unit|tenant)\s*[:\-]\s*([A-Za-z0-9][A-Za-z0-9 _/-]*)', text, re.IGNORECASE)
                    if match:
                        unit = match.group(1).strip()
                        break
                header_index = next((index for index, row in enumerate(rows[:30]) if any(any(token in str(value or '').casefold() for token in ('meter awal', 'meter akhir', 'usage', 'tagihan', 'pemakaian')) for value in row)), None)
                if header_index is None:
                    continue
                headers = [str(value or '').strip() or f'Kolom {index + 1}' for index, value in enumerate(rows[header_index])]
                unit_index = next((index for index, header in enumerate(headers) if header.casefold() in {'unit', 'unit/area', 'tenant'}), None)
                for row in rows[header_index + 1:]:
                    values = {headers[index]: value for index, value in enumerate(row) if index < len(headers) and value not in (None, '')}
                    if not values:
                        continue
                    month = ''
                    for value in row[:4]:
                        text_value = str(value or '').strip()
                        normalized = text_value.casefold()
                        if normalized in month_names:
                            month = month_map.get(normalized, text_value.title())
                            break
                        date_match = re.match(r'^(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})$', text_value)
                        if date_match:
                            month = datetime.strptime(text_value.replace('/', '-'), '%d-%m-%Y').strftime('%B')
                            month = month_map.get(month.casefold(), month)
                            break
                    if file_month and not month:
                        month = file_month
                    if not month:
                        # Some templates put the month in a merged cell or use
                        # a localized label; retain non-empty rows with the
                        # first value as a month-like label.
                        first = str(row[0] or '').strip() if row else ''
                        if first and len(first) <= 24 and not re.search(r'^(total|subtotal|discount|pajak|ppn)', first, re.IGNORECASE):
                            month = first.title()
                        else:
                            continue
                    row_unit = str(row[unit_index]).strip() if unit_index is not None and unit_index < len(row) and row[unit_index] not in (None, '') else unit
                    if not row_unit:
                        row_unit = sheet.title.strip()
                    if not re.search(r'[A-Za-z]', row_unit) or len(row_unit) > 30:
                        continue
                    db.session.add(TenantUtility(unit=row_unit, month=month.title(), source_filename=secure_filename(upload.filename), values_json=json.dumps(values, ensure_ascii=False, default=str)))
                    imported += 1
            db.session.commit()
            return jsonify(success=True, imported=imported)
        except Exception as exc:
            db.session.rollback()
            return jsonify(error=f'Gagal membaca file Excel: {exc}'), 422
    units = [value for (value,) in db.session.query(TenantUtility.unit).distinct().order_by(TenantUtility.unit).all()]
    months = [value for (value,) in db.session.query(TenantUtility.month).distinct().all()]
    selected_units = request.args.getlist('unit')
    selected_months = request.args.getlist('month')
    query = TenantUtility.query
    if selected_units:
        query = query.filter(TenantUtility.unit.in_(selected_units))
    if selected_months:
        query = query.filter(TenantUtility.month.in_(selected_months))
    rows = []
    if selected_units and selected_months:
        for row in query.order_by(TenantUtility.unit, TenantUtility.id).all():
            rows.append({'id': row.id, 'unit': row.unit, 'month': row.month, 'values': json.loads(row.values_json), 'source': row.source_filename})
    return jsonify(data=rows, units=units, months=months, success=True)


@app.route('/api/pdf-upload', methods=['GET', 'POST'])
def pdf_upload_api():
    if request.method == 'POST':
        try:
            if not csrf_valid():
                return jsonify(error='Form kedaluwarsa. Muat ulang halaman.'), 400
            if 'pdf_file' not in request.files:
                return jsonify(error='Tidak ada file yang diupload'), 400
            
            file = request.files['pdf_file']
            if file.filename == '':
                return jsonify(error='Tidak ada file yang dipilih'), 400
            
            if not file.filename.lower().endswith('.pdf'):
                return jsonify(error='Hanya file PDF yang diperbolehkan'), 400
            
            pdf_content = file.read()
            pdf_reader = PdfReader(io.BytesIO(pdf_content))
            extracted_pages = [(page, page.extract_text() or '') for page in pdf_reader.pages]
            worksheet_pages = [(page, content) for page, content in extracted_pages
                               if 'WORKSHEET PURCHASE REQUEST' in content.upper()]
            selected_pages = worksheet_pages[:1] if worksheet_pages else extracted_pages
            plain_text = '\n'.join(content for _, content in selected_pages)
            try:
                layout_text = '\n'.join(page.extract_text(extraction_mode='layout') or '' for page, _ in selected_pages)
            except (TypeError, ValueError):
                layout_text = ''
            positioned_text = '\n'.join(positioned_pdf_text(page) for page, _ in selected_pages)
            if not plain_text.strip() and not layout_text.strip() and not positioned_text.strip():
                return jsonify(error='PDF tidak memiliki teks yang dapat dibaca. PDF hasil scan memerlukan OCR.'), 422
            parsed_data = parse_pdf_text(layout_text)
            plain_data = parse_pdf_text(plain_text)
            positioned_data = parse_pdf_text(positioned_text)
            for field in PDF_FIELDS:
                if field == 'status':
                    continue
                if not parsed_data.get(field):
                    parsed_data[field] = plain_data.get(field, '')
                if not parsed_data.get(field):
                    parsed_data[field] = positioned_data.get(field, '')
            missing = [field for field in ('tanggal', 'no_wpr', 'deskripsi', 'type_budget') if not parsed_data[field]]
            return jsonify(data=parsed_data, missing=missing, success=True)
            
        except Exception as e:
            return jsonify(error=f'Gagal memproses PDF: {str(e)}'), 500
    else:
        try:
            uploads = PdfUpload.query.order_by(PdfUpload.created_at.desc()).all()
            def date_key(row):
                match = re.match(r'^(\d{1,2})[-/]?(\d{1,2})[-/]?(\d{2,4})$', (row.tanggal or '').strip())
                if not match:
                    return (0, 0, 0, '')
                day, month, year = (int(value) for value in match.groups())
                if year < 100:
                    year += 2000
                return (year, month, day, (row.no_wpr or '').upper())
            uploads.sort(key=date_key, reverse=True)
            data = [pdf_upload_dict(u) for u in uploads]
            return jsonify(data=data, success=True)
        except Exception as e:
            return jsonify(error=str(e)), 500


PDF_FIELDS = ('tanggal', 'no_wpr', 'deskripsi', 'type_budget', 'status', 'keterangan')
PDF_STATUSES = ('Cancel', 'Rejected', 'SPK Created', 'PO Created')


def pdf_upload_dict(row):
    return {'id': row.id, 'has_pdf': bool(row.pdf_filename),
            'pdf_filename': row.pdf_filename or '',
            **{field: getattr(row, field) or '' for field in PDF_FIELDS}}


def uploaded_pdf(required=False):
    file = request.files.get('pdf_file')
    if not file or not file.filename:
        if required:
            raise ValueError('Pilih file PDF sebelum menyimpan data.')
        return None
    filename = secure_filename(file.filename)
    if not filename.lower().endswith('.pdf'):
        raise ValueError('Hanya file PDF yang diperbolehkan.')
    content = file.read()
    if not content.startswith(b'%PDF-'):
        raise ValueError('File yang dipilih bukan PDF yang valid.')
    try:
        PdfReader(io.BytesIO(content))
    except Exception as exc:
        raise ValueError('File PDF rusak atau tidak dapat dibaca.') from exc
    # Rewrite pages through PyPDF2 to apply stream compression and discard
    # unnecessary incremental/object overhead before storing in SQLite.
    try:
        reader = PdfReader(io.BytesIO(content))
        writer = PdfWriter()
        for page in reader.pages:
            page.compress_content_streams()
            writer.add_page(page)
        compressed = io.BytesIO()
        writer.write(compressed)
        if compressed.tell() < len(content):
            content = compressed.getvalue()
    except Exception:
        pass
    return filename, content


def pdf_form_data():
    data = {field: request.form.get(field, '').strip() for field in PDF_FIELDS}
    if any(not data[field] or data[field] == '—' for field in ('tanggal', 'no_wpr', 'deskripsi', 'type_budget')):
        raise ValueError('Tanggal, No WPR, deskripsi, dan type budget wajib diisi.')
    if data['status'] and data['status'] not in PDF_STATUSES:
        raise ValueError('Status tidak valid.')
    return data


@app.post('/api/pdf-records')
def pdf_record_create():
    if not csrf_valid():
        return jsonify(error='Form kedaluwarsa. Muat ulang halaman.'), 400
    try:
        pdf = uploaded_pdf(required=True)
        row = PdfUpload(**pdf_form_data(), pdf_filename=pdf[0], pdf_content=pdf[1])
    except ValueError as exc:
        return jsonify(error=str(exc)), 400
    db.session.add(row)
    db.session.commit()
    return jsonify(data=pdf_upload_dict(row), success=True), 201


@app.route('/api/pdf-records/<int:record_id>', methods=['PUT', 'DELETE'])
@app.post('/api/pdf-records/<int:record_id>/edit')
@app.post('/api/pdf-records/<int:record_id>/delete')
def pdf_record_change(record_id):
    if not csrf_valid():
        return jsonify(error='Form kedaluwarsa. Muat ulang halaman.'), 400
    row = db.session.get(PdfUpload, record_id)
    if row is None:
        return jsonify(error='Data tidak ditemukan.'), 404
    if request.method == 'DELETE' or request.path.endswith('/delete'):
        db.session.delete(row)
        db.session.commit()
        return jsonify(success=True)
    try:
        data = pdf_form_data()
        pdf = uploaded_pdf()
    except ValueError as exc:
        return jsonify(error=str(exc)), 400
    for field, value in data.items():
        setattr(row, field, value)
    if pdf:
        row.pdf_filename, row.pdf_content = pdf
    db.session.commit()
    return jsonify(data=pdf_upload_dict(row), success=True)


@app.get('/api/pdf-records/<int:record_id>/pdf')
def pdf_record_file(record_id):
    row = db.session.get(PdfUpload, record_id)
    if row is None or not row.pdf_filename or not row.pdf_content:
        return jsonify(error='File PDF tidak tersedia untuk data ini.'), 404
    return send_file(io.BytesIO(row.pdf_content), mimetype='application/pdf',
                     as_attachment=False, download_name=row.pdf_filename)


PDF_LABELS = {
    'tanggal': r'(?:tanggal\s*permintaan|tanggal\s*wpr|tanggal|date)',
    'no_wpr': r'(?:no\.?\s*wpr|nomor\s*wpr|wpr\s*no\.?|no\.?|nomor)',
    'deskripsi': r'(?:deskripsi(?:\s*pekerjaan)?|uraian(?:\s*pekerjaan)?|pekerjaan|description)',
    'type_budget': r'(?:type\s*budget|tipe\s*budget|jenis\s*budget|budget\s*type|budget)',
    'keterangan': r'(?:keterangan|catatan|remarks|notes?)',
}
PDF_LABEL_PATTERN = re.compile(
    r'^\s*(?:' + '|'.join(PDF_LABELS.values()) + r')\s*(?::|=|\s+-\s|\s{2,}|$)', re.IGNORECASE
)


def positioned_pdf_text(page):
    fragments = []

    def collect(fragment, cm, tm, font, size):
        value = ' '.join(fragment.split())
        if value:
            x, y = float(tm[4]), float(tm[5])
            if cm and len(cm) >= 6:
                x, y = (cm[0] * x + cm[2] * y + cm[4],
                        cm[1] * x + cm[3] * y + cm[5])
            fragments.append((y, x, value))

    try:
        page.extract_text(visitor_text=collect)
    except (TypeError, ValueError, IndexError):
        return ''
    fragments.sort(key=lambda item: (-item[0], item[1]))
    rows = []
    for y, x, value in fragments:
        if not rows or abs(rows[-1][0] - y) > 4:
            rows.append((y, [value]))
        else:
            rows[-1][1].append(value)
    return '\n'.join(' '.join(values) for _, values in rows)


def parse_pdf_text(text):
    text = text.replace('\u00a0', ' ').replace('\u2013', '-').replace('\u2014', '-')
    lines = [' '.join(line.split()) for line in text.splitlines()]
    lines = [line for line in lines if line]
    data = {field: '' for field in PDF_FIELDS if field != 'status'}

    # The worksheet places the description immediately before "Budget Capex/Opex",
    # without a "Deskripsi" label. Keep it separate from the item name above it.
    worksheet_budget = re.search(
        r'\bBudge(?:t)?\s*:?[ \t\r\n]*(.+?)(?=\s+Periode\b|\s+Keterangan\b|\s+Catatan\b|\s+Subtotal\b|$)',
        text, re.IGNORECASE | re.DOTALL
    )
    if worksheet_budget:
        value = ' '.join(worksheet_budget.group(1).split()).strip(' :-')
        standard = re.search(r'\b(Capex|Opex)\b', value, re.IGNORECASE)
        if standard and value.lower().startswith('periode'):
            value = standard.group(1).title()
        if value:
            data['type_budget'] = value
    else:
        budget_label = re.search(r'\bBudge(?:t)?\b', text, re.IGNORECASE)
        if budget_label:
            nearby = text[budget_label.end():budget_label.end() + 80]
            standard = re.search(r'\b(Capex|Opex)\b', nearby, re.IGNORECASE)
            if standard:
                data['type_budget'] = standard.group(1).title()
            value = re.search(r'\b([A-Za-z][A-Za-z /()_-]{1,50}?)(?=\s+Periode\b|\s+Keterangan\b|\s+Catatan\b|\s+Subtotal\b|$)', nearby, re.IGNORECASE)
            if value and not data['type_budget']:
                data['type_budget'] = ' '.join(value.group(1).split()).strip(' :-')
    for index, line in enumerate(lines):
        budget_in_line = re.search(r'\bBudge(?:t)?\b', line, re.IGNORECASE)
        if not budget_in_line or data['deskripsi']:
            continue
        prefix = line[:budget_in_line.start()].strip(' :-')
        if not prefix and index:
            prefix = lines[index - 1].strip(' :-')
        # Item descriptions may legitimately contain terms such as "No.1";
        # only discard prefixes that are clearly table/header labels.
        if len(prefix) >= 3 and not re.match(r'^(?:Qty|Pcs|Kode Barang|Nama Barang|Tipe|Type|Keterangan|Tanggal|Nomor)\b', prefix, re.IGNORECASE):
            data['deskripsi'] = prefix

    worksheet_date = re.search(
        r'Tanggal\s*Permintaan\s*:?\s*(\d{1,2}\s*[-/]\s*\d{1,2}\s*[-/]\s*\d{2,4})',
        text, re.IGNORECASE
    )
    if worksheet_date:
        data['tanggal'] = re.sub(r'\s*([- /])\s*', r'\1', worksheet_date.group(1))
    else:
        date_label = re.search(r'Tanggal\s*Permintaan', text, re.IGNORECASE)
        if date_label:
            nearby = text[date_label.end():date_label.end() + 180]
            value = re.search(r'\b(\d{1,2}\s*[-/]\s*\d{1,2}\s*[-/]\s*\d{2,4})\b', nearby)
            if value:
                data['tanggal'] = re.sub(r'\s*([- /])\s*', r'\1', value.group(1))

    for index, line in enumerate(lines):
        for field, label in PDF_LABELS.items():
            match = re.match(r'^\s*(?:' + label + r')\s*(?::|=|\s+-\s|\s{2,}|$)\s*(.*)$', line, re.IGNORECASE)
            if not match or data[field]:
                continue
            value = match.group(1).strip()
            if not value and index + 1 < len(lines) and not PDF_LABEL_PATTERN.match(lines[index + 1]):
                value = lines[index + 1]
            if field in ('deskripsi', 'keterangan') and value:
                continuation = []
                for next_line in lines[index + 1:index + 4]:
                    if PDF_LABEL_PATTERN.match(next_line):
                        break
                    if next_line != value:
                        continuation.append(next_line)
                if continuation:
                    value = ' '.join([value] + continuation)
            if field == 'tanggal':
                date_match = re.search(r'\b(?:\d{1,2}\s*[-/]\s*\d{1,2}\s*[-/]\s*\d{2,4}|\d{4}-\d{2}-\d{2})\b', value)
                value = date_match.group(0) if date_match else ''
                value = re.sub(r'\s*([- /])\s*', r'\1', value)
            elif field == 'no_wpr':
                wpr_match = re.search(r'\bWPR[A-Za-z0-9/_-]*\b', value, re.IGNORECASE)
                value = wpr_match.group(0) if wpr_match else ''
            elif field == 'keterangan' and re.search(r'\b(?:Qty|Satuan|Spesifikasi|Item Price|Extended Price)\b', value, re.IGNORECASE):
                value = ''
            data[field] = value

    if not data['no_wpr']:
        match = re.search(r'\bWPR(?:[/-][A-Za-z0-9]+)+\b', text, re.IGNORECASE)
        if match:
            data['no_wpr'] = match.group(0)
    return data


@app.errorhandler(ReportError)
def handle_report_error(error):
    return jsonify(error=str(error)), error.status


@app.after_request
def response_headers(response):
    response.headers['Cache-Control'] = 'public, max-age=3600' if request.path.startswith('/static/') else 'no-store'
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['Referrer-Policy'] = 'no-referrer'
    return response


# Register additional page routes
from supabase_viewer import register_supabase_viewer
register_supabase_viewer(app)

from berita_acara_routes import register_berita_acara_routes
register_berita_acara_routes(app, db, upload_folder='uploads')

with app.app_context():
    db.create_all()
    with db.engine.begin() as connection:
        columns = {row[1] for row in connection.exec_driver_sql('PRAGMA table_info(pdf_uploads)')}
        if 'pdf_filename' not in columns:
            connection.execute(sql_text('ALTER TABLE pdf_uploads ADD COLUMN pdf_filename VARCHAR(255)'))
        if 'pdf_content' not in columns:
            connection.execute(sql_text('ALTER TABLE pdf_uploads ADD COLUMN pdf_content BLOB'))
        if 'status' not in columns:
            connection.execute(sql_text('ALTER TABLE pdf_uploads ADD COLUMN status VARCHAR(30)'))


if __name__ == '__main__':
    app.run(host='0.0.0.0', port=int(os.getenv('PORT', '8080')), debug=False)
