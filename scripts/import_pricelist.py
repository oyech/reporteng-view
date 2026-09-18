"""Import the supplied FLT workbook without modifying its values."""
import argparse
import json
from pathlib import Path
from openpyxl import load_workbook


def extract(path):
    workbook = load_workbook(path, data_only=True)
    prices, criteria, detail = [workbook[name] for name in ('Lampiran 1', 'Lampiran 2', 'Lampiran 3')]
    rates = []
    for row in (7, 9, 11, 13, 15):
        for col in (2, 5, 8, 11, 14):
            rates.append(dict(code=prices.cell(row, col).value, category=prices.cell(row, 1).value,
                              type=prices.cell(5, col).value,
                              prices=[prices.cell(row + 1, col + n).value for n in range(3)],
                              criteria=[criteria.cell(row + 2, col + n).value for n in range(3)], examples=[]))
    by_code = {rate['code']: rate for rate in rates}
    codes = {}
    for row in criteria.iter_rows(min_row=35, values_only=True):
        for col in (1, 4, 7, 10, 13):
            value = row[col]
            if value in by_code:
                codes[col] = value
            elif value is not None:
                by_code[codes[col]]['examples'].append(str(value).strip())
    notes = {label: {criteria.cell(row + n, 1).value: criteria.cell(row + n, 2).value for n in range(3)}
             for label, row in [('Waktu', 21), ('Tingkat kesulitan', 26), ('Safety', 31)]}
    entries = []
    category = group = parent = ''
    for number, cells in enumerate(detail.iter_rows(min_row=8, max_col=20, values_only=True), 8):
        cells = [value.strip() if isinstance(value, str) else value for value in cells]
        if isinstance(cells[1], str) and cells[2]:
            category, group, parent = cells[2], '', ''
            continue
        name = cells[2] or cells[3]
        if not name:
            continue
        free = category.startswith('Pekerjaan Yang Tidak Dikenakan Biaya')
        has_detail = any(value is not None for value in cells[4:17]) or free
        if cells[2]:
            if isinstance(cells[1], (int, float)):
                group, parent = name, ''
            else:
                parent = name
        if not has_detail:
            continue
        entries.append(dict(row=number, category=category, group=group,
                            parent=parent if parent != name else '', name=name,
                            note=cells[4], code=cells[5], volume=cells[6], unit=cells[7],
                            service_type=cells[9], service=cells[13], material=cells[14],
                            total=cells[15], rounded=cells[16], free=free, source_values=cells))
    workbook.close()
    return dict(source=Path(path).name, period='April 2019–Maret 2020',
                entries=entries, rates=rates, criteria_notes=notes)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook')
    parser.add_argument('--output', type=Path, default=Path(__file__).resolve().parents[1] / 'data' / 'pricelist.json')
    args = parser.parse_args()
    data = extract(args.workbook)
    args.output.write_text(json.dumps(data, ensure_ascii=True, indent=2) + '\n', encoding='utf-8')
    print(f"Imported {len(data['entries'])} entries and {len(data['rates'])} service rates.")
